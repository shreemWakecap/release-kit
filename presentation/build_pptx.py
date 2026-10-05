#!/usr/bin/env python3
"""build_pptx.py: slides.json (from extract.js) -> connected-environment-1.0.pptx
Native shapes, editable text boxes (Arial, exact line spacing), pictures, speaker notes, real slide titles.
Run with the venv: .venv/bin/python build_pptx.py
"""
import json, os, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "connected-environment-1.0.pptx")
PX = 9525  # EMU per CSS px at 96 dpi
SLACK = 8  # px of extra width on text boxes so PowerPoint never wraps earlier than the browser did

data = json.load(open(os.path.join(HERE, "slides.json")))
prs = Presentation()
prs.slide_width = Emu(data["width"] * PX)
prs.slide_height = Emu(data["height"] * PX)
cp = prs.core_properties
cp.title, cp.author, cp.subject = "Connected Environment 1.0", "WakeCap", "Release presentation, October 2026 (internal)"
cp.keywords, cp.comments = "Connected Environment, Weather Station, Lightning, Gas", "Live production screens, 4 October 2026. Speaker notes on every slide."

# theme font -> Arial so text added later matches
try:
    for rel in prs.slide_master.part.rels.values():
        if rel.reltype == RT.THEME:
            part = rel.target_part
            part._blob = part.blob.replace(b"Calibri", b"Arial")
except Exception as e:  # cosmetic only
    print("theme font patch skipped:", e)

layout = prs.slide_layouts[5]  # Title Only


def emu(v):
    return Emu(int(round(v * PX)))


def rgb(hexs):
    return RGBColor.from_string(hexs)


def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def add_rect(slide, x, y, w, h, fill=None, line=None, radius=0, dash=False):
    """fill: (hex, alpha) | None ; line: (hex, width_px) | None"""
    m = min(w, h)
    if radius <= 0:
        kind, adj = MSO_SHAPE.RECTANGLE, None
    elif abs(w - h) < 1.5 and radius >= m / 2 - 0.5:
        kind, adj = MSO_SHAPE.OVAL, None
    else:
        kind, adj = MSO_SHAPE.ROUNDED_RECTANGLE, min(0.5, radius / m)
    shp = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
    strip_style(shp)
    if adj is not None:
        shp.adjustments[0] = adj
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill[0])
        if fill[1] < 0.999:
            clr = shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = etree.SubElement(clr, qn("a:alpha"))
            a.set("val", str(int(fill[1] * 100000)))
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line[0])
        shp.line.width = emu(line[1])
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    else:
        shp.line.fill.background()
    return shp


def add_shape_item(slide, it):
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    fill = (it["fill"]["hex"], it["fill"]["a"]) if it["fill"] else None
    b = it["border"]
    if not b:
        add_rect(slide, x, y, w, h, fill, None, it["radius"])
        return
    bw = b["w"]
    widths = {bw["t"], bw["r"], bw["b"], bw["l"]}
    if len(widths) == 1:
        t = bw["t"]
        add_rect(slide, x + t / 2, y + t / 2, w - t, h - t, fill, (b["hex"], t), max(0, it["radius"] - t / 2), b["dash"])
    else:
        if fill:
            add_rect(slide, x, y, w, h, fill, None, it["radius"])
        for side, rect in (("l", (x, y, bw["l"], h)), ("r", (x + w - bw["r"], y, bw["r"], h)),
                           ("t", (x, y, w, bw["t"])), ("b", (x, y + h - bw["b"], w, bw["b"]))):
            if bw[side] > 0:
                add_rect(slide, *rect, fill=(b["hex"], 1), line=None)


def add_text_item(slide, it, title_shape=None):
    x, w = it["x"], it["w"]
    if it["align"] == "left":
        w += SLACK
    elif it["align"] == "center":
        x -= SLACK / 2
        w += SLACK
    else:
        x -= SLACK
        w += SLACK
    if title_shape is not None:
        shp = title_shape
        shp.left, shp.top, shp.width, shp.height = emu(x), emu(it["y"]), emu(w), emu(it["h"])
    else:
        shp = slide.shapes.add_textbox(emu(x), emu(it["y"]), emu(w), emu(it["h"]))
    tf = shp.text_frame
    tf.word_wrap = not it["nowrap"]
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[it["align"]]
    p.line_spacing = Pt(it["lineHeight"] * 0.75)
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    for r in it["runs"]:
        if r.get("br"):
            p.add_line_break()
            continue
        run = p.add_run()
        run.text = r["text"]
        f = run.font
        f.name = "Arial"
        f.size = Pt(r["size"] * 0.75)
        f.bold = bool(r["bold"])
        f.italic = bool(r["italic"])
        f.color.rgb = rgb(r["color"])
        if r.get("spacing"):
            run._r.get_or_add_rPr().set("spc", str(int(round(r["spacing"] * 0.75 * 100))))
    return shp


for idx, sd in enumerate(data["slides"], 1):
    slide = prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(sd["bg"])
    title_ph = slide.shapes.title
    title_used = False
    for it in sd["items"]:
        if it["type"] == "shape":
            add_shape_item(slide, it)
        elif it["type"] == "image":
            slide.shapes.add_picture(os.path.join(HERE, it["src"]), emu(it["x"]), emu(it["y"]), emu(it["w"]), emu(it["h"]))
        elif it["type"] == "text":
            if it["tag"] == "H1" and not title_used:
                add_text_item(slide, it, title_ph)
                title_used = True
            else:
                add_text_item(slide, it)
    if not title_used and title_ph is not None:
        title_ph._element.getparent().remove(title_ph._element)
    if sd["notes"].strip():
        slide.notes_slide.notes_text_frame.text = sd["notes"].strip()

prs.save(OUT)
print("wrote", OUT, "|", len(prs.slides), "slides |", round(os.path.getsize(OUT) / 1024), "KB")
